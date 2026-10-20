"""PDF and Word export for the Sheet module."""
import io
from django.utils import timezone

HEADERS=["S.No","Date","Product Name","HSN Code","Vendor","Stock","Sales","Balance","MRP (Rs.)","Disc %","Price (Rs.)"]

def _n(v,places=None):
    if v is None: return "-"
    if places is not None: return f"{v:,.{places}f}"
    s=f"{v:f}"
    return s.rstrip("0").rstrip(".") if "." in s else s

def rows_for(entries):
    return [[str(e.sno),e.entry_date.strftime("%d-%m-%Y"),e.product.name,e.product.hsn_code,e.vendor.name if e.vendor_id else "-",
             _n(e.stock),_n(e.sales),_n(e.balance),_n(e.mrp,2),(_n(e.disc)+"%") if e.disc is not None else "-",_n(e.price,2)] for e in entries]

def subtitle(start,end):
    rng=f"{start or 'start'} to {end or 'today'}" if (start or end) else "All dates"
    return f"{rng}  |  Generated {timezone.localtime().strftime('%d-%m-%Y %H:%M')}"

def build_pdf(entries,shop_name,start=None,end=None):
    from .pdf_reports import build_table_pdf
    entries=list(entries); rows=rows_for(entries)
    st=sum((e.stock for e in entries),0); sa=sum(((e.sales or 0) for e in entries),0)
    return build_table_pdf(shop_name=shop_name,title="Sheet - Date-wise Stock",subtitle=subtitle(start,end),
        headers=HEADERS,rows=rows,widths_mm=[13,23,58,22,40,17,17,19,22,22,25],right_from=5,
        summary=[("Entries",str(len(entries))),("Stock",_n(st) if entries else "0"),("Sales",_n(sa) if entries else "0"),("Balance",_n(st-sa) if entries else "0")],
        empty_text="No entries found")

def build_docx(entries,shop_name,start=None,end=None):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Mm, RGBColor
    d=Document()
    sec=d.sections[0]; sec.orientation=WD_ORIENT.LANDSCAPE
    sec.page_width,sec.page_height=Mm(297),Mm(210)
    for m in ("left_margin","right_margin","top_margin","bottom_margin"): setattr(sec,m,Mm(14))
    d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10)
    d.add_heading(f"{shop_name} - Sheet",level=1)
    p=d.add_paragraph(subtitle(start,end)); p.runs[0].font.size=Pt(9); p.runs[0].font.color.rgb=RGBColor(0x66,0x76,0x84)
    rows=rows_for(entries)
    t=d.add_table(rows=1,cols=len(HEADERS)); t.style="Table Grid"
    def shade(cell,hexcol):
        tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement("w:shd")
        sh.set(qn("w:val"),"clear"); sh.set(qn("w:color"),"auto"); sh.set(qn("w:fill"),hexcol); tcPr.append(sh)
    for i,h in enumerate(HEADERS):
        c=t.rows[0].cells[i]; c.text=""; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); shade(c,"0E7C7B")
        if i>=5: c.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    # repeat header row on each page
    trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement("w:tblHeader"); th.set(qn("w:val"),"true"); trPr.append(th)
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=v
            if i>=5: cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    if not rows:
        cells=t.add_row().cells; cells[2].text="No entries found"
    widths=[Mm(11),Mm(21),Mm(52),Mm(20),Mm(36),Mm(16),Mm(16),Mm(18),Mm(20),Mm(20),Mm(22)]
    for row in t.rows:
        for i,w in enumerate(widths): row.cells[i].width=w
    buf=io.BytesIO(); d.save(buf); return buf.getvalue()
